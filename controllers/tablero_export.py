import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generar_tablero_excel(datos_resumen, nombre_archivo="Tablero_Control.xlsx"):
  wb = openpyxl.Workbook()

  # Hoja 1: Dashboard / Tablero
  ws_dashboard = wb.active
  ws_dashboard.title = "Tablero"
  ws_dashboard.views.sheetView[0].showGridLines = True

  # Estilos corporativos
  font_titulo = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
  font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
  fill_header = PatternFill(
      start_color="1F4E78", end_color="1F4E78", fill_type="solid"
  )
  fill_titulo = PatternFill(
      start_color="2F5597", end_color="2F5597", fill_type="solid"
  )
  borde_delgado = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )

  # Título del Tablero
  ws_dashboard.merge_cells("A1:E2")
  celda_titulo = ws_dashboard["A1"]
  celda_titulo.value = "TABLERO DE CONTROL DE CLIENTES Y TRAMAS"
  celda_titulo.font = font_titulo
  celda_titulo.fill = fill_titulo
  celda_titulo.alignment = Alignment(horizontal="center", vertical="center")

  # Cabeceras de métricas / resumen
  headers = [
      "Métrica / Indicador",
      "Total Registros",
      "Monto Total Remuneración",
      "Promedio",
      "Estado",
  ]
  ws_dashboard.row_dimensions[4].height = 25

  for col_idx, header in enumerate(headers, start=1):
    cell = ws_dashboard.cell(row=4, column=col_idx, value=header)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = borde_delgado

  # Llenado de datos simulados o provistos para el tablero
  fila_inicio = 5
  for i, item in enumerate(
      datos_resumen or [("Sede Principal", 15, 153750.00, 10250.00, "Activo")]
  ):
    row = fila_inicio + i
    ws_dashboard.cell(row=row, column=1, value=item[0])
    ws_dashboard.cell(row=row, column=2, value=item[1])
    celda_monto = ws_dashboard.cell(row=row, column=3, value=item[2])
    celda_monto.number_format = '"S/ "#,##0.00'
    celda_prom = ws_dashboard.cell(row=row, column=4, value=item[3])
    celda_prom.number_format = '"S/ "#,##0.00'
    ws_dashboard.cell(row=row, column=5, value=item[4])

    for col_idx in range(1, 6):
      c = ws_dashboard.cell(row=row, column=col_idx)
      c.border = borde_delgado
      c.alignment = Alignment(
          horizontal="left" if col_idx == 1 else "center", vertical="center"
      )

  # Autoajustar columnas
  for col in ws_dashboard.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_dashboard.column_dimensions[col_letter].width = max(max_len + 3, 15)

  wb.save(nombre_archivo)
  print(f"Tablero generado con éxito: {nombre_archivo}")


